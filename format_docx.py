# -*- coding: utf-8 -*-
"""
论文格式一键排版

命令行用法：
    python format_docx.py 我的论文.docx
    python format_docx.py 我的论文.docx -o 排好版.docx -r rules.yaml

图形界面用法：
    python gui.py
"""

import argparse
from collections import Counter
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt

# 中文号数 → 磅值，rules.yaml 里也可以直接写数字
CN_SIZE = {
    "初号": 42, "小初": 36, "一号": 26, "小一": 24,
    "二号": 22, "小二": 18, "三号": 16, "小三": 15,
    "四号": 14, "小四": 12, "五号": 10.5, "小五": 9,
}

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

# rules.yaml 里哪一段对应 Word 里的哪个样式名
# 注意：一个 Word 样式只能被配置一次，重复映射会互相覆盖
STYLE_MAP = {
    "body": ["Normal", "正文"],
    "h1": ["Heading 1", "标题 1"],
    "h2": ["Heading 2", "标题 2"],
    "h3": ["Heading 3", "标题 3"],
    "caption": ["Caption", "题注"],
    "reference": ["Bibliography", "参考文献"],   # 文档里没有这个样式就跳过
}


def to_pt(value):
    """把 12 / '小四' / '12pt' 统一转成磅值"""
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("pt", "").replace("磅", "")
    if text in CN_SIZE:
        return CN_SIZE[text]
    return float(text)


def set_line_spacing(pf, value):
    """小于等于 5 当作倍数，大于 5 当作固定磅值"""
    if isinstance(value, (int, float)) and value <= 5:
        pf.line_spacing = float(value)
    else:
        pf.line_spacing = Pt(to_pt(value))


def set_style_format(style, cfg):
    """把一份配置应用到某个 Word 样式上"""
    # ---------- 字体 ----------
    font_cn = cfg.get("font_cn", "宋体")
    font_en = cfg.get("font_en", "Times New Roman")
    style.font.name = font_en                      # 西文
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_cn)          # 中文（必须单独设）
    rfonts.set(qn("w:ascii"), font_en)
    rfonts.set(qn("w:hAnsi"), font_en)

    # ---------- 字号与粗细 ----------
    style.font.size = Pt(to_pt(cfg.get("size", 12)))
    style.font.bold = bool(cfg.get("bold", False))

    # ---------- 段落 ----------
    pf = style.paragraph_format
    pf.alignment = ALIGN.get(str(cfg.get("align", "left")).lower(),
                             WD_ALIGN_PARAGRAPH.LEFT)
    set_line_spacing(pf, cfg.get("line_spacing", 1.5))
    pf.space_before = Pt(float(cfg.get("space_before", 0)))
    pf.space_after = Pt(float(cfg.get("space_after", 0)))

    # ---------- 首行缩进（按字符，Word 原生支持） ----------
    chars = float(cfg.get("first_line_chars", 0))
    ppr = style.element.get_or_add_pPr()
    ind = ppr.get_or_add_ind()
    size_pt = to_pt(cfg.get("size", 12))
    if chars > 0:
        ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
        ind.set(qn("w:firstLine"), str(int(chars * size_pt * 20)))
    else:
        ind.set(qn("w:firstLineChars"), "0")
        ind.set(qn("w:firstLine"), "0")


def add_page_number(section, align="center"):
    """在页脚插入 PAGE 域。页脚已有文字时不覆盖，只记一笔。"""
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()

    if p.text.strip():
        return f"页脚已有内容「{p.text.strip()[:16]}」，跳过页码"

    p.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.CENTER)

    for run in list(p.runs):
        run._r.getparent().remove(run._r)

    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    for node in (begin, instr, sep, placeholder, end):
        run._r.append(node)
    return None


def apply_rules(doc, rules):
    """把规则应用到已打开的文档对象，返回统计结果。"""
    notes = []

    # ---------- 1. 页面 ----------
    page = rules.get("page", {})
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(float(page.get("margin_top", 2.5)))
        section.bottom_margin = Cm(float(page.get("margin_bottom", 2.5)))
        section.left_margin = Cm(float(page.get("margin_left", 3.0)))
        section.right_margin = Cm(float(page.get("margin_right", 2.5)))
        if page.get("page_number", True):
            msg = add_page_number(section, page.get("page_number_align", "center"))
            if msg:
                notes.append(msg)

    # ---------- 2. 样式 ----------
    applied, missing = [], []
    for key, names in STYLE_MAP.items():
        if key not in rules:
            continue
        cfg = rules[key]
        target = None
        for name in names:
            try:
                target = doc.styles[name]
                break
            except KeyError:
                continue
        if target is None:
            missing.append(key)
            continue
        set_style_format(target, cfg)
        applied.append(target.name)

    # ---------- 3. 统计 ----------
    counts = Counter(p.style.name for p in doc.paragraphs)
    h1 = counts.get("Heading 1", 0) + counts.get("标题 1", 0)
    h2 = counts.get("Heading 2", 0) + counts.get("标题 2", 0)
    body = counts.get("Normal", 0) + counts.get("正文", 0)

    return {
        "applied": list(dict.fromkeys(applied)),
        "missing": missing,
        "h1": h1,
        "h2": h2,
        "body": body,
        "notes": notes,
    }


def process(src, out, rules):
    """读文件 → 应用规则 → 保存。返回统计结果。"""
    src, out = Path(src), Path(out)
    if not src.exists():
        raise FileNotFoundError(f"找不到文件：{src}")
    if src.suffix.lower() != ".docx":
        raise ValueError("只支持 .docx 文件。如果是 .doc，先在 Word 里另存为 .docx。")

    doc = Document(str(src))
    stats = apply_rules(doc, rules)
    doc.save(str(out))
    return stats


def load_rules(path):
    return yaml.safe_load(Path(path).read_text(encoding="utf-8"))


def main():
    ap = argparse.ArgumentParser(description="按规则文件一键排版 Word 论文")
    ap.add_argument("input", help="要排版的 .docx 文件")
    ap.add_argument("-o", "--output", help="输出文件名，默认在原名后加 -已排版")
    ap.add_argument("-r", "--rules", default="rules.yaml", help="规则文件")
    args = ap.parse_args()

    src = Path(args.input)
    out = Path(args.output) if args.output else src.with_name(src.stem + "-已排版.docx")
    rules = load_rules(args.rules)

    stats = process(src, out, rules)

    print("=" * 52)
    print("排版完成")
    print("=" * 52)
    print(f"已应用样式：{'、'.join(stats['applied'])}")
    if stats["missing"]:
        print(f"文档里没有这些样式，已跳过：{'、'.join(stats['missing'])}")
    for n in stats["notes"]:
        print(f"  · {n}")
    print()
    print(f"一级标题段落：{stats['h1']}")
    print(f"二级标题段落：{stats['h2']}")
    print(f"正文段落：    {stats['body']}")
    print()

    if stats["h1"] == 0:
        print("⚠ 没有检测到任何「标题 1」段落。")
        print("  说明你的论文是用加粗放大手动做的标题，不是 Word 的标题样式。")
        print("  这种情况下本工具只能改正文，标题层级改不了。")
        print("  解决：在 Word 里选中标题 → 开始 → 样式 → 标题 1（或按 Ctrl+Alt+1）。")
        print()

    print(f"已保存：{out}")


if __name__ == "__main__":
    main()
