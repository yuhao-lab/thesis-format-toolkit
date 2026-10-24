# -*- coding: utf-8 -*-
"""
生成一份「格式很乱」的测试论文，用来验证排版工具。

用法：
    python make_test_doc.py

会输出 测试论文-原始.docx，正文故意用宋体五号、单倍行距、无缩进，
跑一遍 format_docx.py 之后对比，就能看出工具改了什么。
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

doc = Document()

# 故意把正文设成五号、单倍行距、无缩进、两端不对齐
normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.0
normal.paragraph_format.first_line_indent = Pt(0)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.add_heading("第一章  绪论", level=1)

doc.add_paragraph(
    "新能源汽车产业在最近五年里完成了从政策驱动到市场驱动的转变。"
    "这种转变带来的直接后果是，整车电子电气架构的复杂度大幅上升，"
    "线束数量、控制器数量和信号数量都比传统燃油车高出一个量级。"
)

doc.add_heading("1.1  研究背景", level=2)

doc.add_paragraph(
    "智能座舱作为用户感知最强的一环，其语音交互质量直接影响整车评价。"
    "然而在实际项目中，语音功能的测试往往依赖人工逐条点检，"
    "缺乏统一的用例标准和可复现的判定口径。"
)

doc.add_paragraph(
    "本文以某车型座舱语音系统为对象，尝试建立一套可复用的测试用例集，"
    "并对常见的失效模式进行分类归因。"
)

doc.add_heading("1.2  研究内容", level=2)

doc.add_paragraph(
    "主要工作包括三部分：用例设计、执行记录、失效归因。"
)

doc.add_heading("第二章  系统设计", level=1)

doc.add_heading("2.1  整体架构", level=2)

doc.add_paragraph(
    "系统由车机端、云端和测试台架三部分组成。车机端负责采集语音输入，"
    "云端完成识别与语义解析，测试台架用于复现问题场景。"
)

cap = doc.add_paragraph("图 2-1  系统整体架构")
cap.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.save("测试论文-原始.docx")
print("已生成：测试论文-原始.docx")
